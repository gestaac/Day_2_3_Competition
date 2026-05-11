import openpyxl
wb = openpyxl.load_workbook(r'C:\Users\CHESTER BAUTISTA\Day_2_3_Competition\WSA2025_54_Cyber_Security_marking_scheme_RevisedScore.xlsx', data_only=True)
out = []
for sh in wb.sheetnames:
    ws = wb[sh]
    out.append(f"===SHEET: {sh} (rows={ws.max_row} cols={ws.max_column})===")
    for row in ws.iter_rows(values_only=True):
        if any(c is not None and str(c).strip() != "" for c in row):
            out.append("\t".join("" if c is None else str(c) for c in row))
with open(r'C:\Users\CHESTER BAUTISTA\Day_2_3_Competition\_xlsx_extract.txt','w',encoding='utf-8') as f:
    f.write("\n".join(out))
print("DONE", len(out))
